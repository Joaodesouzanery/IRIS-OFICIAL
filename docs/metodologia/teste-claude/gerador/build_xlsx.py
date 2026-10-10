import os
HERE=os.path.dirname(os.path.abspath(__file__))
OUTDIR=os.path.dirname(HERE)
import sys, urllib.parse
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.formatting.rule import ColorScaleRule, CellIsRule, FormulaRule
from openpyxl.workbook.defined_name import DefinedName
from openpyxl.utils import get_column_letter as L
from model import *
OUT=sys.argv[1] if len(sys.argv)>1 else os.path.join(OUTDIR,'IMQN_agencias_workbook.xlsx'); DEMO=len(sys.argv)>2  # DEMO: respostas sintéticas só para teste de fórmula
dims,crits,conds=load()
wb=Workbook()
HF=PatternFill('solid',fgColor='1F3A5F'); HFONT=Font(bold=True,color='FFFFFF'); WR=Alignment(wrap_text=True,vertical='top')
INP=PatternFill('solid',fgColor='FFF8DC'); thin=Side(style='thin',color='D0D0D0'); BD=Border(left=thin,right=thin,top=thin,bottom=thin)
def head(ws,row,cols,widths=None):
    for i,c in enumerate(cols,1):
        x=ws.cell(row,i,c); x.fill=HF; x.font=HFONT; x.alignment=Alignment(wrap_text=True,vertical='center'); x.border=BD
    if widths:
        for i,w in enumerate(widths,1): ws.column_dimensions[L(i)].width=w
    ws.freeze_panes=ws.cell(row+1,1)
def put(ws,r,c,v,**k):
    x=ws.cell(r,c,v); x.alignment=WR; x.border=BD
    for a,b in k.items(): setattr(x,a,b)
    return x
# ---------- LEIA-ME
ws=wb.active; ws.title='LEIA-ME'; ws.column_dimensions['A'].width=120
txt=[('IMQN — Qualidade Regulatória das 12 agências federais · workbook de trabalho (rev2022)',True),
('Base: planilha "Matriz de Avaliação da Maturidade da Qualidade Normativa" (programa INFRA Competitividade, rev2022). Estrutura reescrita aqui de forma própria.',False),('',False),
('COMO USAR',True),
('1. Aba "Avaliação": uma linha por agência × condição (12 × 73 = 876). Preencha só as células amarelas: Resposta (SIM / NÃO / ?), URL da evidência, data da consulta, status.',False),
('2. Use a aba "Roteiro" para abrir uma busca direcionada ao portal da agência, por critério. O portal raiz e a lei de criação estão em "Agências".',False),
('3. A aba "Cálculo" aplica a regra de nível e a aba "Ranking" soma tudo. Nada precisa ser digitado nelas.',False),('',False),
('REGRA DE NÍVEL (decisão metodológica deste workbook — a planilha original não a define)',True),
('• Níveis cumulativos: Inexistente (0) → Inicial (0,35) → Gerenciado (0,70) → Melhoria Contínua (1,00). Valores editáveis em "Parâmetros".',False),
('• Para cada critério: o nível-base é o mais alto cujas condições estão TODAS atendidas (e todas as dos níveis abaixo). A fração do nível seguinte é a proporção das condições dele atendidas.',False),
('• Exemplo: Inicial completo + 2 de 4 condições do Gerenciado → 0,35 + 0,35 × 0,5 = 0,525. Condição de nível acima de uma lacuna não conta (maturidade não pula degrau).',False),
('• Nota da dimensão = Σ (peso do critério na dimensão × classificação) × peso da dimensão. IMQN = Σ das 6 dimensões (0–100). Pesos: AIR 25 · Estoque 20 · Participação Social 15 · Agenda 15 · ARR 15 · Processo 10.',False),('',False),
('DUAS NOTAS, NUNCA UMA SÓ',True),
('• ESTIMADO: conta toda resposta SIM. • COMPROVADO: conta só SIM com evidência (URL) E status "validado". Nenhuma coleta automática nasce validada.',False),
('• Teto verificável por fonte pública: aba "Condições" mostra quais condições só dado interno comprova (ex.: Capacitação é 100% interna). Ver "Ranking" para o teto.',False),('',False),
('GUARDRAILS',True),
('• Avalia a INSTITUIÇÃO, nunca pessoas. Não registre CPF, telefone/e-mail pessoal, nomes de servidores ou juízo reputacional individual. Capacitação entra só como percentual agregado.',False),
('• Resposta em branco = "sem avaliação", NÃO é "não atende". Sem evidência, não se atribui nota.',False),
('• Classificação pública/interna, perguntas e termos de busca são PROPOSTA deste workbook; base legal marcada "a conferir" até revisão jurídica.',False),('',False),
('ABAS: Parâmetros · Dimensões · Critérios · Condições · Agências · Roteiro · Fontes · Sinais · Mapa · Monitor · Achados · Avaliação · Cálculo · Ranking · Auditoria',False),('Fontes/Sinais/Mapa/Monitor: de onde vem cada dado e como monitorá-lo (fonte única: fontes-monitoramento.json). Nada nelas atribui nota.',False)]
for i,(t,b) in enumerate(txt,1):
    x=ws.cell(i,1,t); x.alignment=Alignment(wrap_text=True,vertical='top'); x.font=Font(bold=b,size=14 if i==1 else 11)
if DEMO: ws.cell(len(txt)+2,1,'⚠ VERSÃO DE TESTE com respostas sintéticas — NÃO usar').font=Font(bold=True,color='FF0000')
# ---------- Parâmetros
ws=wb.create_sheet('Parâmetros'); head(ws,1,['Nível','Valor','Nome'],[22,12,30])
for i,(n,v,nm) in enumerate([('Inicial',0.35,'v_ini'),('Gerenciado',0.7,'v_ger'),('Melhoria Contínua',1.0,'v_mc')],2):
    put(ws,i,1,n); put(ws,i,2,v,fill=INP); put(ws,i,3,nm)
    wb.defined_names[nm]=DefinedName(nm,attr_text=f"'Parâmetros'!$B${i}")
put(ws,6,1,'Inexistente = 0 (fixo). Peso dos critérios e dimensões nas abas Critérios/Dimensões.')
# ---------- Dimensões
ws=wb.create_sheet('Dimensões'); head(ws,1,['ID','Dimensão','Peso (pts)','O que mede','Base legal','Base legal conferida?'],[5,36,10,60,50,16])
MED={1:'A agência avalia impacto e alternativas antes de criar norma?',2:'Consultas e audiências, com devolutiva às contribuições?',3:'Acervo normativo levantado, indexado e consolidado?',4:'Planejamento público do que será regulado — e cumprido?',5:'Fluxo de criação de normas padronizado e monitorado?',6:'Normas em vigor são reavaliadas depois de aplicadas?'}
for i,d in enumerate(dims,2):
    put(ws,i,1,d['id']); put(ws,i,2,d['nome']); put(ws,i,3,d['peso']); put(ws,i,4,MED[d['id']]); put(ws,i,5,BASE[d['criterios'][0]['codigo'].split('_')[0]]); put(ws,i,6,'Não',fill=INP)
put(ws,8,2,'TOTAL',font=Font(bold=True)); put(ws,8,3,'=SUM(C2:C7)',font=Font(bold=True))
# ---------- Critérios
ws=wb.create_sheet('Critérios'); head(ws,1,['Código','Critério','Dim.','Peso na dimensão','Condições Inicial','Condições Gerenciado','Condições Melh. Contínua','Total','Termos de busca','Onde procurar'],[10,34,6,12,10,12,12,8,48,52])
for i,c in enumerate(crits,2):
    put(ws,i,1,c['codigo']); put(ws,i,2,c['nome']); put(ws,i,3,c['dim']); put(ws,i,4,c['peso'])
    for j,n in enumerate(['Inicial','Gerenciado','Melhoria Contínua'],5): put(ws,i,j,f"=COUNTIFS(Condições!$C$2:$C$74,$A{i},Condições!$E$2:$E$74,\"{n}\")")
    put(ws,i,8,f'=SUM(E{i}:G{i})'); put(ws,i,9,c['keywords']); put(ws,i,10,c['onde'])
put(ws,12,2,'Conferência: Σ pesos dentro de cada dimensão deve ser 1',font=Font(italic=True))
# ---------- Condições
ws=wb.create_sheet('Condições'); cols=['ID','Critério','Dim.','Nível','Nível (nome)','Ordem','Condição (texto da matriz, corrigido)','Pergunta sim/não','Verificabilidade (proposta)','Tipo de evidência','Por quê','Base legal','Texto original']
head(ws,1,cols,[16,10,5,10,17,6,60,60,14,24,40,36,50])
# coluna B = critério, C=critério? keep fixed order: A id, B crit, C = crit (dup for COUNTIFS) -> reorganize
ws.delete_cols(1,15)
cols=['ID','Dim.','Critério','Nível-chave','Nível','Ordem','Condição (texto da matriz, corrigido)','Pergunta sim/não','Verificabilidade (proposta)','Tipo de evidência','Por quê','Base legal','Texto original','Origem na planilha (aba!célula)','Autoria']
head(ws,1,cols,[16,5,10,10,17,6,60,60,14,24,40,36,50,18,48])
for i,c in enumerate(conds,2):
    vals=[c['id'],c['dim'],c['crit'],c['nivel'],NIVLBL[c['nivel']],c['ordem'],c['texto'],c['pergunta'],c['verif'],c['tipo'],c['motivo'],c['base'],c['original'],c['celula'],'Condição = planilha INFRA rev2022 (texto oficial). Pergunta, verificabilidade e tipo de evidência = proposta deste workbook.']
    for j,v in enumerate(vals,1): put(ws,i,j,v)
# Critérios usa col C (critério) e col E (nível nome) -> ok
# ---------- Agências
ws=wb.create_sheet('Agências'); head(ws,1,['Sigla','Agência','Setor','Lei de criação','Ministério supervisor','Portal (raiz)','Acesso do portal medido em 2026-10-09','Observação'],[9,46,28,28,24,34,22,60])
OBS={'ANPD':'Natureza jurídica/enquadramento na Lei 13.848/2019 a conferir antes de comparar com as demais.','ANA':'Regulação de saneamento (Lei 14.026/2020) inclui normas de referência — verificar agenda própria.','ANCINE':'Criada por MP 2.228-1/2001; verificar agenda regulatória bienal publicada.'}
for i,(s,n,st,lei,slug) in enumerate(AG,2):
    a=ACESSO[slug]; put(ws,i,1,s,font=Font(bold=True)); put(ws,i,2,n); put(ws,i,3,st); put(ws,i,4,lei); put(ws,i,5,'a conferir',fill=INP)
    x=put(ws,i,6,f'https://www.gov.br/{slug}/pt-br'); x.hyperlink=f'https://www.gov.br/{slug}/pt-br'; x.font=Font(color='0563C1',underline='single')
    put(ws,i,7,{200:'HTTP 200',403:'bloqueado (403) neste ambiente — conferir no navegador',0:'sem resposta — conferir no navegador'}[a]); put(ws,i,8,OBS.get(s,''))
# ---------- Roteiro (busca por agência × critério)
ws=wb.create_sheet('Roteiro'); head(ws,1,['Agência']+[c['codigo'] for c in crits],[10]+[16]*10)
for i,(s,n,st,lei,slug) in enumerate(AG,2):
    put(ws,i,1,s,font=Font(bold=True))
    for j,c in enumerate(crits,2):
        term=c['keywords'].split(';')[0].strip()
        url='https://www.google.com/search?q='+urllib.parse.quote(f'site:gov.br/{slug} "{term}"')
        x=put(ws,i,j,'buscar'); x.hyperlink=url; x.font=Font(color='0563C1',underline='single')
put(ws,14,1,'Cada link abre uma busca restrita ao portal da agência pelo 1º termo do critério (ver aba Critérios para os demais termos).')

import json as _j
REG=_j.load(open(os.path.join(OUTDIR,'fontes-monitoramento.json'),encoding='utf-8'))
def link(ws,r,c,url,text=None):
    x=put(ws,r,c,text or url)
    if url.startswith('http'): x.hyperlink=url; x.font=Font(color='0563C1',underline='single')
    return x
# ---------- Fontes (nacionais)
ws=wb.create_sheet('Fontes'); head(ws,1,['ID','Fonte','Entidade','URL','Formato','Periodicidade','Status da URL','Cobre','Para que serve neste monitoramento','Risco de acesso/automação'],[6,40,18,46,22,14,28,16,60,34])
for i,f in enumerate(REG['nacionais'],2):
    put(ws,i,1,f['id']); put(ws,i,2,f['nome']); put(ws,i,3,f['entidade']); link(ws,i,4,f['url']); put(ws,i,5,f['formato']); put(ws,i,6,f['periodicidade']); put(ws,i,7,f['status']); put(ws,i,8,f['cobre']); put(ws,i,9,f['uso']); put(ws,i,10,f['risco'])
# ---------- Sinais
ws=wb.create_sheet('Sinais'); head(ws,1,['ID','Sinal observável','Automação (A/B/C)','Como detectar','Fontes','Frequência'],[6,52,12,64,34,12])
for i,s in enumerate(REG['sinais'],2):
    for j,v in enumerate([s['id'],s['nome'],s['automacao'],s['metodo'],s['fontes'],s['frequencia']],1): put(ws,i,j,v)
put(ws,len(REG['sinais'])+3,1,'A = página/DOU/dataset detecta sozinho · B = exige ler PDF/painel (extração + revisão humana) · C = dado interno ou leitura humana (LAI). Nenhum sinal vira nota: vira SUGESTÃO de condição, sempre "pendente".')
ws.merge_cells(start_row=len(REG['sinais'])+3,start_column=1,end_row=len(REG['sinais'])+3,end_column=6)
# ---------- Mapa
ws=wb.create_sheet('Mapa'); head(ws,1,['Condição ID','Critério','Nível','Ordem','Origem na planilha','Condição','Sinais que a evidenciam','Automação','Fontes primárias dos sinais','Frequência de checagem'],[16,10,17,6,16,60,16,10,50,16])
SID={s['id']:s for s in REG['sinais']}
for i,m_ in enumerate(REG['mapa'],2):
    ss=m_['sinais']
    for j,v in enumerate([m_['condicao'],m_['critério'],NIVLBL[m_['nivel']],m_['ordem'],m_['celula_planilha'],m_['texto'],', '.join(ss),m_['grau'],' | '.join(dict.fromkeys(SID[s]['fontes'] for s in ss)),' / '.join(dict.fromkeys(SID[s]['frequencia'] for s in ss))],1): put(ws,i,j,v)
ws.auto_filter.ref=f'A1:J{len(REG["mapa"])+1}'
# resumo A/B/C por critério
r0=len(REG['mapa'])+4
put(ws,r0,1,'Resumo de automação por critério',font=Font(bold=True)); head_cols=['Critério','A','B','C','Total']
for j,v in enumerate(head_cols,1):
    x=ws.cell(r0+1,j,v); x.fill=HF; x.font=HFONT
for i,c in enumerate(crits,r0+2):
    put(ws,i,1,c['codigo'])
    for j,g in enumerate('ABC',2): put(ws,i,j,f'=COUNTIFS($B$2:$B${len(REG["mapa"])+1},$A{i},$H$2:$H${len(REG["mapa"])+1},"{g}")')
    put(ws,i,5,f'=SUM(B{i}:D{i})')
put(ws,r0+12,1,'TOTAL',font=Font(bold=True))
for j in range(2,6): put(ws,r0+12,j,f'=SUM({L(j)}{r0+2}:{L(j)}{r0+11})',font=Font(bold=True))
# ---------- Monitor
ws=wb.create_sheet('Monitor'); head(ws,1,['Agência','Critério','Chave','URL a monitorar','Status da URL','Plataforma de participação','Métrica própria da agência','Painel público','Termo para redescobrir','Verificada por navegador?'],[9,10,7,70,34,28,40,28,36,14])
r=2
for m_ in REG['monitor']:
    for f in m_['fontes']:
        put(ws,r,1,m_['agencia']); put(ws,r,2,m_['criterio']); put(ws,r,3,m_['chave'])
        if f['url'].startswith('http'): link(ws,r,4,f['url'])
        else: put(ws,r,4,f['url'])
        put(ws,r,5,f['status']); put(ws,r,6,m_['plataforma_ps']); put(ws,r,7,m_['metrica_propria']); put(ws,r,8,m_['painel'])
        q=m_['busca']+' '+m_['agencia']; x=put(ws,r,9,q); x.hyperlink='https://www.google.com/search?q='+urllib.parse.quote(q+' site:gov.br'); x.font=Font(color='0563C1',underline='single')
        put(ws,r,10,'Não',fill=INP); r+=1
ws.auto_filter.ref=f'A1:J{r-1}'
ws.conditional_formatting.add(f'E2:E{r-1}',FormulaRule(formula=['LEFT(E2,11)="a descobrir"'],fill=PatternFill('solid',bgColor='FFC7CE')))
ws.conditional_formatting.add(f'E2:E{r-1}',FormulaRule(formula=['LEFT(E2,6)="padrão"'],fill=PatternFill('solid',bgColor='FFEB9C')))
ws.conditional_formatting.add(f'E2:E{r-1}',FormulaRule(formula=['LEFT(E2,9)="encontrad"'],fill=PatternFill('solid',bgColor='C6EFCE')))
# ---------- Achados
ws=wb.create_sheet('Achados'); head(ws,1,['#','Achado (pesquisa de 09/out/2026)','Implicação para os critérios / monitoramento','Fonte'],[4,70,60,50])
ACH=[('O Decreto 10.139/2019 (citado na planilha rev2022 para o Estoque) foi substituído pelo Decreto 12.002/2024 (22/04/2024), que mantém a consolidação normativa obrigatória e traz a revogação de atos exauridos no art. 64.','A base legal da dimensão Estoque está desatualizada na própria matriz. Atualizar a citação e conferir os artigos antes de qualquer publicação.','gov.br/anatel (consolidação); gov.br/anac (estoque); gov.br/aneel (notícia 2024)'),
('A plataforma federal de participação mudou de "Participa + Brasil" para "Brasil Participativo"; ANATEL, ANA, ANTT e ANEEL ainda usam sistemas próprios.','Não existe fonte única de consultas: o monitor precisa de adaptador por plataforma (5 famílias) e do dataset aberto da ANEEL/ANAC quando houver.','páginas de participação das agências'),
('Várias agências já publicam métricas próprias de qualidade: IQR (ANATEL, piloto 72,5%), IQAIR (ANA, média 89 em 2023), ICAR (ANTT), ICPN (ANEEL), Selo de Boas Práticas (ANM), RAPS (ANVISA).','Aproveitar como evidência e como régua: o critério pode passar de "existe processo de avaliação da qualidade das AIR" para um indicador numérico comparável.','gov.br/anatel; gov.br/ana (IQAIR); gov.br/antt; gov.br/aneel; gov.br/anm'),
('ANVISA fechou a Agenda 2024-2025 com 27% de conclusão (meta 50%) e 84% das normas ligadas à agenda (meta 75%); ANATEL reporta 42,3% de execução (set/2025).','Cumprimento de agenda e aderência norma×agenda JÁ são publicados em % — dá para exigir limiares numéricos (≥50% / ≥80%) em vez de "existe relatório".','gov.br/anvisa (informe de resultados); gov.br/anatel (relatório de execução)'),
('ANS: não localizei a Agenda 2026-2028 aprovada (a anterior era 2023-2025). A página de AIR da ANS tem o slug "copy_of_".','Risco de falso negativo no monitor (página vigente pode ter outro endereço) → status "a descobrir" até verificação humana.','gov.br/ans'),
('ANCINE não realizou consultas públicas em 2024 (relatório de ouvidoria); ANPD tem consulta da Agenda 2027-2028 aberta até 16/10/2026; ANM tem Audiência Pública 02/2026 aberta até 14/10/2026.','Ausência de consulta num ano não é, por si, falha — o sinal precisa de denominador (normas que exigiam consulta). Eventos com prazo geram alerta de calendário.','relatório de ouvidoria ANCINE; gov.br/anpd; gov.br/anm'),
('8 de 12 portais gov.br falharam no acesso direto neste ambiente (7 com HTTP 403 e 1 sem resposta).','O coletor pode sofrer bloqueio semelhante em produção: tratar 403 como "não sei" (já é o desenho do coletor atual), nunca como "não publica".','teste de acesso de 09/out/2026'),
('Estoque, Processo Normativo e Capacitação: sem página própria localizada em ANP, ANM, ANCINE, ANPD (e Processo em ANVISA e ANPD).','São os pontos com menos fonte pública: 6 linhas do Monitor ficaram "a descobrir"; Capacitação só por LAI.','Aba Monitor'),
('ANPD aparece com AIR e agenda próprias, mas o enquadramento na Lei 13.848/2019 como agência reguladora federal deve ser conferido.','Comparabilidade do ranking: considerar rotular a ANPD à parte até a conferência jurídica.','gov.br/anpd')]
for i,a in enumerate(ACH,2):
    put(ws,i,1,i-1)
    for j,v in enumerate(a,2): put(ws,i,j,v)

# ---------- Avaliação
ws=wb.create_sheet('Avaliação'); head(ws,1,['Agência','Condição ID','Critério','Dim.','Nível','Ordem','Condição','Resposta (SIM/NÃO/?)','URL da evidência','Data da consulta','Status revisão','Observação (sem dados pessoais)','Conta estimado','Conta comprovado'],[9,16,10,5,17,6,60,12,40,14,14,40,10,10])
r=2
import random; random.seed(7)
for (s,*_) in AG:
    for c in conds:
        put(ws,r,1,s); put(ws,r,2,c['id']); put(ws,r,3,c['crit']); put(ws,r,4,c['dim']); put(ws,r,5,NIVLBL[c['nivel']]); put(ws,r,6,c['ordem']); put(ws,r,7,c['texto'])
        resp=None
        if DEMO and s=='ANTT': resp='SIM' if (c['nivel']=='inicial') or (c['crit']=='AGE' and c['nivel']=='gerenciado' and c['ordem'] in('I','II')) else None
        if DEMO and s=='ANM': resp='SIM'
        if DEMO and s=='ANEEL': resp='SIM' if c['crit']=='EST' and c['nivel']=='inicial' else 'NÃO'
        put(ws,r,8,resp,fill=INP); put(ws,r,9,'https://exemplo.gov.br/doc' if DEMO and s in('ANM',) else None,fill=INP); put(ws,r,10,None,fill=INP)
        put(ws,r,11,('validado' if DEMO and s=='ANM' and c['crit'] in('AIR_MET','PS') else 'pendente'),fill=INP); put(ws,r,12,None,fill=INP)
        put(ws,r,13,f'=IF(H{r}="SIM",1,0)'); put(ws,r,14,f'=IF(AND(H{r}="SIM",K{r}="validado",I{r}<>""),1,0)')
        r+=1
N=r-1
dv=DataValidation(type='list',formula1='"SIM,NÃO,?"',allow_blank=True); ws.add_data_validation(dv); dv.add(f'H2:H{N}')
dv2=DataValidation(type='list',formula1='"pendente,preliminar,em_revisao,validado,rejeitado"',allow_blank=False); ws.add_data_validation(dv2); dv2.add(f'K2:K{N}')
ws.conditional_formatting.add(f'H2:H{N}',CellIsRule(operator='equal',formula=['"SIM"'],fill=PatternFill('solid',bgColor='C6EFCE')))
ws.conditional_formatting.add(f'H2:H{N}',CellIsRule(operator='equal',formula=['"NÃO"'],fill=PatternFill('solid',bgColor='FFC7CE')))
ws.auto_filter.ref=f'A1:N{N}'
# ---------- Cálculo
ws=wb.create_sheet('Cálculo'); cols=['Agência','Critério','Dim.','Peso na dim.','n Inicial','n Gerenc.','n MC','ok Inicial (est)','ok Gerenc. (est)','ok MC (est)','ok Inicial (comp)','ok Gerenc. (comp)','ok MC (comp)','Classificação ESTIMADA (0–1)','Classificação COMPROVADA (0–1)']
head(ws,1,cols,[9,10,5,10,8,8,8,10,10,10,10,10,10,16,16])
rr=2
rngA=f'Avaliação!$A$2:$A${N}'; rngC=f'Avaliação!$C$2:$C${N}'; rngE=f'Avaliação!$E$2:$E${N}'
for (s,*_) in AG:
    for c in crits:
        put(ws,rr,1,s); put(ws,rr,2,c['codigo']); put(ws,rr,3,c['dim']); put(ws,rr,4,c['peso'])
        for j,n in enumerate(['Inicial','Gerenciado','Melhoria Contínua'],5): put(ws,rr,j,f'=COUNTIFS(Condições!$C$2:$C$74,$B{rr},Condições!$E$2:$E$74,"{n}")')
        for j,n in enumerate(['Inicial','Gerenciado','Melhoria Contínua'],8): put(ws,rr,j,f'=COUNTIFS({rngA},$A{rr},{rngC},$B{rr},{rngE},"{n}",Avaliação!$M$2:$M${N},1)')
        for j,n in enumerate(['Inicial','Gerenciado','Melhoria Contínua'],11): put(ws,rr,j,f'=COUNTIFS({rngA},$A{rr},{rngC},$B{rr},{rngE},"{n}",Avaliação!$N$2:$N${N},1)')
        for col,(a,b,c3) in ((14,('H','I','J')),(15,('K','L','M'))):
            put(ws,rr,col,f'=IF({a}{rr}<E{rr},v_ini*{a}{rr}/E{rr},IF({b}{rr}<F{rr},v_ini+(v_ger-v_ini)*{b}{rr}/F{rr},IF({c3}{rr}<G{rr},v_ger+(v_mc-v_ger)*{c3}{rr}/G{rr},v_mc)))',number_format='0.000')
        rr+=1
M=rr-1
# ---------- Ranking
ws=wb.create_sheet('Ranking'); cols=['Agência','Condições respondidas','Cobertura da avaliação']+[f"{d['nome'].split(' (')[0]} (máx {d['peso']})" for d in dims]+['IMQN ESTIMADO (0–100)','IMQN COMPROVADO (0–100)','Posição (estimado)','Cobertura comprovada (condições)']
head(ws,1,cols,[10,12,12,16,16,16,16,16,16,14,14,10,14])
ws.row_dimensions[1].height=48
for i,(s,*_) in enumerate(AG,2):
    put(ws,i,1,s,font=Font(bold=True))
    put(ws,i,2,f'=COUNTIFS(Avaliação!$A$2:$A${N},$A{i},Avaliação!$H$2:$H${N},"SIM")+COUNTIFS(Avaliação!$A$2:$A${N},$A{i},Avaliação!$H$2:$H${N},"NÃO")')
    put(ws,i,3,f'=B{i}/{len(conds)}',number_format='0%')
    for d in dims:
        col=3+d['id']
        put(ws,i,col,f"=SUMPRODUCT((Cálculo!$A$2:$A${M}=$A{i})*(Cálculo!$C$2:$C${M}={d['id']})*Cálculo!$D$2:$D${M}*Cálculo!$N$2:$N${M})*Dimensões!$C${d['id']+1}",number_format='0.0')
    put(ws,i,10,f'=SUM(D{i}:I{i})',number_format='0.0',font=Font(bold=True))
    put(ws,i,11,f'=SUMPRODUCT(Cálculo!$D$2:$D${M}*(Cálculo!$A$2:$A${M}=$A{i})*Cálculo!$O$2:$O${M}*SUMIF(Dimensões!$A$2:$A$7,Cálculo!$C$2:$C${M},Dimensões!$C$2:$C$7))',number_format='0.0',font=Font(bold=True))
    put(ws,i,12,f'=IF(B{i}=0,"—",RANK(J{i},$J$2:$J$13))')
    put(ws,i,13,f'=SUMIFS(Avaliação!$N$2:$N${N},Avaliação!$A$2:$A${N},$A{i})/{len(conds)}',number_format='0%')
ws.conditional_formatting.add('D2:I13',ColorScaleRule(start_type='min',start_color='F8696B',mid_type='percentile',mid_value=50,mid_color='FFEB84',end_type='max',end_color='63BE7B'))
ws.conditional_formatting.add('J2:K13',ColorScaleRule(start_type='num',start_value=0,start_color='F8696B',mid_type='num',mid_value=50,mid_color='FFEB84',end_type='num',end_value=100,end_color='63BE7B'))
put(ws,15,1,'Notas: (1) "Posição" só aparece após a agência ter ≥1 condição respondida. (2) ESTIMADO soma todo SIM; COMPROVADO exige SIM + URL + status validado. (3) Em branco ≠ NÃO.',font=Font(italic=True))
ws.merge_cells('A15:M15')
# teto verificável só por fonte pública (cascata em Python)
def casc(vals):
    ini,ger,mc=vals
    if ini[0]<ini[1]: return .35*ini[0]/ini[1]
    if ger[0]<ger[1]: return .35+.35*ger[0]/ger[1]
    if mc[0]<mc[1]: return .7+.3*mc[0]/mc[1]
    return 1.0
teto=0
for d in dims:
    for c in d['criterios']:
        v=[]
        for niv,_,_ in NIV:
            cs=[x for x in conds if x['crit']==c['codigo'] and x['nivel']==niv]
            v.append((sum(1 for x in cs if x['verif']=='pública'),len(cs)))
        teto+=casc(v)*c['peso']*d['peso']
put(ws,16,1,f'Teto verificável só por fonte pública (proposta de classificação): {teto:.2f} de 100 — o restante exige dado interno do órgão.',font=Font(bold=True)); ws.merge_cells('A16:M16')
print('TETO',round(teto,2))
# ---------- Auditoria
ws=wb.create_sheet('Auditoria'); head(ws,1,['#','Onde (planilha original)','Achado','Impacto','O que este workbook faz'],[4,28,60,40,50])
AUD=[('INDICADOR!C5:L11','O nível é escolhido digitando 1/0 em 4 células por critério (one-hot); nada liga as condições das abas AIR/PS/… ao nível.','Classificação é juízo livre, não reprodutível nem auditável.','Resposta por condição + regra de cascata em fórmula.'),
('INDICADOR!E18:E23','Fórmulas com referências fixas (M4*C3*C5+…); AIR e ARR repetem a conta, as outras 4 dimensões usam outra forma.','Difícil estender ou auditar; erro de célula passa despercebido.','Uma única fórmula por critério (Cálculo), parâmetros nomeados (v_ini, v_ger, v_mc).'),
('Abas AIR/PS/Estoque/Agenda/Processo/ARR','Nº de condições por nível é assimétrico (Processo: 1 em Gerenciado; Estoque: 4 em Gerenciado e 2 em MC; Capacitação AIR: 10).','Cada condição pesa diferente conforme o critério; evoluir de nível "custa" mais em uns que em outros.','Mantido (fidelidade à fonte) e exposto em Critérios. Decisão metodológica pendente: normalizar ou não.'),
('Todas','Nível "Inexistente" não tem condições (é ausência do Inicial) e a regra de cascata não está escrita.','Dois avaliadores podem enquadrar o mesmo órgão em níveis diferentes.','Cascata documentada em LEIA-ME e implementada em fórmula.'),
('INDICADOR!M4:M10','Escala 0 / 0,35 / 0,70 / 1,00 não é linear e está só implícita nas células.','Mudança de escala exige editar fórmulas.','Parâmetros editáveis e nomeados.'),
('MATRIZ × INDICADOR','Mesmos textos de nível nas duas abas; linhas formatadas até Z1000.','Duplicação; arquivo de 232 KB para 73 condições.','Fonte única na aba Condições.'),
('Textos (vários)','Erros: "equadramento", "intermediario", "de 25% da 49%", "junho de2020", "Participação   Social" (3 espaços).','Ruído em busca/relatórios.','Corrigidos na coluna de trabalho; texto original preservado na última coluna de Condições.'),
('AIR!E3:G7','Tabela solta ("Total MINFRA 102 / básico 61 = 0,60 / intermediário 52 = 0,51 / avançado 0 / ARR 0") sem ano nem fonte.','Dado de capacitação de um órgão sem rastreabilidade.','Não importado. Se for evidência, entra em Avaliação (AIR_CAP) com URL/documento e data.'),
('AIR!A4:C6 (Capacitação)','Faixas inconsistentes entre níveis (ex.: básico 50–74% no Inicial, 75% no Gerenciado; intermediário/avançado misturados) e item IV só em Gerenciado.','Confunde o enquadramento; difícil medir.','Mantidas como estão; sugerido transformar em indicador numérico único (% por nível de capacitação).'),
('Base legal','A planilha não cita artigo por dimensão; PS e ARR exigem conferência (Lei 13.848/2019 arts. 9º–11; Decreto 10.411/2020).','Risco de citação errada em ranking público.','Coluna "Base legal conferida?" = Não até revisão.'),
('Capacitação (AIR e ARR)','São 16 das 73 condições e só dado interno as comprova.','Nota comprovada por fonte pública tem teto abaixo de 100.','Teto exposto na aba Ranking.')]
for i,a in enumerate(AUD,2):
    put(ws,i,1,i-1)
    for j,v in enumerate(a,2): put(ws,i,j,v)
for w in wb.worksheets:
    w.sheet_view.zoomScale=100
wb.calculation.fullCalcOnLoad=True
wb.save(OUT); print('ok',OUT,N,M)
