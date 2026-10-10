import os
HERE=os.path.dirname(os.path.abspath(__file__))
OUTDIR=os.path.dirname(HERE)
import openpyxl, json, re
SRC=os.path.join(os.path.dirname(OUTDIR),'Matriz Qualidade Normativa_rev2022.xlsx')
wb=openpyxl.load_workbook(SRC)
def fix(t):
    t=re.sub(r'\s+',' ',t.replace('\n',' ')).strip()
    for a,b in [('equadramento','enquadramento'),('intermediario','intermediário'),('de 25% da 49%','de 25% a 49%'),('de2020','de 2020'),('Participação Social','Participação Social')]:
        t=t.replace(a,b)
    return t
LV={'A':'melhoria_continua','B':'gerenciado','C':'inicial'}
# sheet -> (dimensao, [(crit_code, crit_name, peso_na_dim)]) ; AIR/ARR have section headers in col A
DIMS=[('AIR',1,'Análise de Impacto Regulatório (AIR)',25,{'Capacitação AIR':('AIR_CAP','Capacitação em AIR',.30),'Metodologia AIR':('AIR_MET','Metodologia de AIR',.35),'Processo AIR':('AIR_PRO','Processo de AIR',.35)}),
('PS',2,'Participação Social',15,None),('Estoque',3,'Gestão de Estoque Regulatório',20,None),('Agenda',4,'Agenda Regulatória',15,None),('Processo',5,'Gestão do Processo Normativo',10,None),
('ARR',6,'Análise de Resultado Regulatório (ARR)',15,{'Capacitação ARR':('ARR_CAP','Capacitação em ARR',.30),'Metodologia ARR':('ARR_MET','Metodologia de ARR',.35),'Processo ARR':('ARR_PRO','Processo de ARR',.35)})]
out=[]
for sh,did,dn,peso,secs in DIMS:
    ws=wb[sh]; cur=None
    crits={}
    if not secs:
        code={'PS':'PS','Estoque':'EST','Agenda':'AGE','Processo':'PRO'}[sh]; crits[code]=dict(codigo=code,nome=dn,peso=1.0,conds=[]); cur=code
    for r in range(2,ws.max_row+1):
        a=ws.cell(r,1).value
        if secs and a in secs:
            c,n,p=secs[a]; crits[c]=dict(codigo=c,nome=n,peso=p,conds=[]); cur=c; continue
        for col,lv in zip('ABC',LV.values()):
            v=ws[f'{col}{r}'].value
            if v and cur:
                m=re.match(r'^\s*([IVX]+)\s*[–-]\s*(.*)$',v.strip(),re.S)
                crits[cur]['conds'].append(dict(nivel=lv,ordem=m.group(1),texto=fix(m.group(2)),original=v.strip(),celula=f'{sh}!{col}{r}'))
    out.append(dict(id=did,aba=sh,nome=dn,peso=peso,criterios=list(crits.values())))
# MATRIZ resumo
mz=wb['MATRIZ']; 
json.dump(out,open(os.path.join(HERE,'matriz.json'),'w'),ensure_ascii=False,indent=1)
for d in out:
    for c in d['criterios']:
        from collections import Counter
        print(d['id'],c['codigo'],c['peso'],dict(Counter(x['nivel'] for x in c['conds'])))
print(sum(len(c['conds']) for d in out for c in d['criterios']))
