"""ANS: le os anexos PowerPoint (.pptx) listados em ans_inventario.json (estado NAO_PDF).
Uso: python3 -I scripts/ans_pptx.py ans_inventario.json fonte/ans texto_ans ans_pptx.json
Baixa para fonte/ans/pptx/, extrai o texto de cada slide com python-pptx (sem LibreOffice) para texto_ans/anexo_pptx_<ref>.txt
e grava ans_pptx.json (sha256, slides, caracteres, nomes de diretores citados, se cita relator/voto)."""
import sys, json, os, re, hashlib, urllib.request
inv_f, fonte, texto, out = sys.argv[1:5]
from pptx import Presentation
inv = json.load(open(inv_f)); os.makedirs(f'{fonte}/pptx', exist_ok=True); os.makedirs(texto, exist_ok=True)
NOMES = ['Damous', 'Medeiros', 'Aquino', 'Secchin', 'Soares', 'Ângelo', 'Celina']
res = []
for p in inv['pdfs']:
    if p['estado'] != 'NAO_PDF' or not p['url'].lower().endswith('.pptx'): continue
    f = f"{fonte}/pptx/{p['ref']}.pptx"
    if not os.path.exists(f):
        req = urllib.request.Request(p['url'], headers={'User-Agent': 'Mozilla/5.0'})
        open(f, 'wb').write(urllib.request.urlopen(req, timeout=90).read())
    b = open(f, 'rb').read(); pr = Presentation(f); linhas = []
    for i, s in enumerate(pr.slides, 1):
        t = []
        for sh in s.shapes:
            if sh.has_text_frame: t.append(' '.join(sh.text_frame.text.split()))
            if getattr(sh, 'has_table', False) and sh.has_table:
                for r in sh.table.rows: t.append(' | '.join(' '.join(c.text.split()) for c in r.cells))
        linhas.append(f'[slide {i}] ' + ' // '.join(x for x in t if x))
    txt = '\n'.join(linhas); tf = f"{texto}/anexo_pptx_{p['ref']}.txt"; open(tf, 'w', encoding='utf-8').write(txt)
    res.append(dict(ref=p['ref'], url=p['url'], arquivo=f, texto=tf, sha256=hashlib.sha256(b).hexdigest(), slides=len(pr.slides), caracteres=len(txt),
                    diretores_citados=[n for n in NOMES if n in txt], cita_relator=bool(re.search(r'relator', txt, re.I)),
                    cita_voto=bool(re.search(r'\bvot(o|ou|aç)', txt, re.I))))
json.dump(res, open(out, 'w'), ensure_ascii=False, indent=1)
print('pptx lidos', len(res), [(r['ref'], r['slides'], r['caracteres'], r['diretores_citados'], r['cita_relator'], r['cita_voto']) for r in res])
